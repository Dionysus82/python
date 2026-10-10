"""Synthetic PDF, Word and Excel writers with a visible synthetic-data watermark."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.domain.banner import FICTIONAL_BANK, SYNTHETIC_BANNER

FIXED = datetime(2026, 10, 10, 0, 0, 0)


def render_pdf(title: str, pages: list[list[str]]) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        invariant=1,
        title=title,
        author=FICTIONAL_BANK,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
    )
    style = ParagraphStyle(
        "Body",
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1C2430"),
    )
    story = []
    for index, paragraphs in enumerate(pages):
        if index:
            story.append(PageBreak())
        for paragraph in paragraphs:
            story.append(Paragraph(escape(paragraph), style))
            story.append(Spacer(1, 4))

    def watermark(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#8C2F2F"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(16 * mm, 8 * mm, SYNTHETIC_BANNER)
        canvas.restoreState()

    document.build(story, onFirstPage=watermark, onLaterPages=watermark)
    return buffer.getvalue()


def _normalise_office(data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            replacement = zipfile.ZipInfo(filename=info.filename, date_time=(2026, 10, 10, 0, 0, 0))
            replacement.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(replacement, payload)
    return buffer.getvalue()


def render_docx(title: str, sections: list[tuple[str, list[str]]]) -> bytes:
    document = Document()
    core = document.core_properties
    core.author = "Harbour Meridian Bank Limited (fictional)"
    core.title = title
    core.created = FIXED
    core.modified = FIXED
    core.comments = SYNTHETIC_BANNER
    document.sections[0].header.paragraphs[0].text = SYNTHETIC_BANNER
    document.sections[0].footer.paragraphs[0].text = SYNTHETIC_BANNER
    document.add_heading(title, level=1)
    document.add_paragraph(SYNTHETIC_BANNER)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return _normalise_office(buffer.getvalue())


def render_xlsx(title: str, sheets: list[dict]) -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    header_fill = PatternFill("solid", fgColor="12263A")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri")
    for spec in sheets:
        sheet = workbook.create_sheet(spec["name"])
        sheet["A1"] = SYNTHETIC_BANNER
        sheet["A1"].font = Font(color="8C2F2F", bold=True, name="Calibri", size=9)
        sheet.merge_cells("A1:G1")
        sheet["A2"] = title
        sheet["A3"] = spec.get("note", "Fictional Harbour Meridian Bank Limited workbook.")
        headers = spec["headers"]
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(4, column, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True)
        for row_index, row in enumerate(spec["rows"], start=5):
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, column, value)
                if isinstance(value, str) and value.startswith("="):
                    cell.number_format = "0.00"
        if spec["rows"]:
            end_column = chr(64 + len(headers))
            end_row = 4 + len(spec["rows"])
            sheet.tables  # ensure attribute exists
            from openpyxl.worksheet.table import Table, TableStyleInfo

            table = Table(displayName=spec.get("table_name", "EvidenceTable"), ref=f"A4:{end_column}{end_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        sheet.oddHeader.left.text = SYNTHETIC_BANNER
        sheet.page_setup.fitToPage = True
        sheet.freeze_panes = "A5"
    workbook.properties.creator = "Harbour Meridian Bank Limited (fictional)"
    workbook.properties.title = title
    workbook.properties.created = FIXED
    workbook.properties.modified = FIXED
    workbook.properties.description = SYNTHETIC_BANNER
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _normalise_office(buffer.getvalue())


def render_csv(rows: list[list[str]]) -> bytes:
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([SYNTHETIC_BANNER])
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def render_txt(paragraphs: list[str]) -> bytes:
    body = SYNTHETIC_BANNER + "\n\n" + "\n\n".join(paragraphs) + "\n\n" + SYNTHETIC_BANNER + "\n"
    return body.encode("utf-8")
