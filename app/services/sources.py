"""Secure source viewer. Filesystem paths are never returned."""

from __future__ import annotations

import csv
import io
import uuid
from typing import NoReturn

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit.service import write_audit
from app.authorisation.service import allowed_customer_ids, visible_document_ids
from app.config import get_settings
from app.domain.banner import SYNTHETIC_BANNER
from app.domain.errors import AccessDeniedError
from app.ingestion.store import TrustedDocumentStore
from app.models.tables import Document, DocumentChunk
from app.security.principal import Principal


def open_source(
    session: Session,
    principal: Principal,
    document_id: str,
    *,
    page: int | None = None,
    sheet: str | None = None,
) -> dict:
    if not principal.can("document.read"):
        raise AccessDeniedError("This role cannot view documents.")
    allowed = allowed_customer_ids(session, principal)
    visible = visible_document_ids(session, principal, allowed)
    if document_id not in visible:
        deny_missing(session, principal, document_id)
    document = session.get(Document, document_id)
    if document is None:
        deny_missing(session, principal, document_id)
    store = TrustedDocumentStore(get_settings().object_store_dir)
    data = store.get(document.storage_key)
    preview = _preview(document, data, page=page, sheet=sheet)
    chunks = session.scalars(select(DocumentChunk).where(DocumentChunk.document_id == document_id)).all()
    write_audit(
        session,
        event_type="document_view",
        outcome="ok",
        detail="Source viewer opened.",
        principal=principal,
        document_id=document_id,
        correlation_id=uuid.uuid4().hex,
        aml_access=document.sensitivity == "aml",
    )
    return {
        "banner": SYNTHETIC_BANNER,
        "document_id": document.document_id,
        "title": document.title,
        "original_filename": document.original_filename,
        "controlled_reference": document.controlled_reference,
        "risk_domain": document.risk_domain,
        "reporting_date": document.reporting_date,
        "version": document.version,
        "superseded": document.superseded,
        "superseded_by_id": document.superseded_by_id,
        "authoring_function": document.authoring_function,
        "confidentiality": document.confidentiality,
        "checksum_sha256": document.checksum_sha256,
        "ingestion_status": document.ingestion_status,
        "malware_scan_status": document.malware_scan_status,
        "page_count": document.page_count,
        "sheet_metadata": document.sheet_metadata,
        "section_metadata": document.section_metadata,
        "prompt_injection_suspected": document.prompt_injection_suspected,
        "incomplete": document.incomplete,
        "viewer_route": document.source_viewer_route,
        "preview": preview,
        "locators": [chunk.locator for chunk in chunks[:40]],
    }


def document_bytes(session: Session, principal: Principal, document_id: str) -> tuple[bytes, str, str]:
    if not principal.can("document.download"):
        raise AccessDeniedError("This role cannot download documents.")
    allowed = allowed_customer_ids(session, principal)
    if document_id not in visible_document_ids(session, principal, allowed):
        deny_missing(session, principal, document_id)
    document = session.get(Document, document_id)
    if document is None:
        deny_missing(session, principal, document_id)
    data = TrustedDocumentStore(get_settings().object_store_dir).get(document.storage_key)
    write_audit(
        session,
        event_type="document_download",
        outcome="ok",
        detail="Authorised download.",
        principal=principal,
        document_id=document_id,
        aml_access=document.sensitivity == "aml",
    )
    return data, document.mime_type, document.safe_filename


def deny_missing(session: Session, principal: Principal, document_id: str) -> NoReturn:
    write_audit(
        session,
        event_type="access_denied",
        outcome="denied",
        detail="Document request blocked before the object was read.",
        principal=principal,
        document_id=None,
        correlation_id=uuid.uuid4().hex,
    )
    raise AccessDeniedError("The requested customer is outside your authorised portfolio or does not exist.")


def _preview(document: Document, data: bytes, *, page: int | None, sheet: str | None) -> dict:
    if document.mime_type == "application/pdf":
        import fitz

        pdf = fitz.open(stream=data, filetype="pdf")
        index = max(1, page or 1)
        index = min(index, pdf.page_count)
        pixmap = pdf[index - 1].get_pixmap(matrix=fitz.Matrix(1.4, 1.4))
        return {"kind": "pdf", "page": index, "page_count": pdf.page_count, "png": pixmap.tobytes("png")}
    if document.mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), data_only=True)
        name = sheet if sheet in book.sheetnames else book.sheetnames[0]
        grid = []
        for row in book[name].iter_rows(max_row=30, max_col=8, values_only=True):
            grid.append(["" if cell is None else str(cell) for cell in row])
        return {"kind": "xlsx", "sheet": name, "sheets": book.sheetnames, "rows": grid}
    text = data.decode("utf-8", errors="replace")
    if document.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document as Docx

        paragraphs = [para.text for para in Docx(io.BytesIO(data)).paragraphs if para.text.strip()]
        return {"kind": "docx", "paragraphs": paragraphs[:40]}
    if document.mime_type == "text/csv":
        rows = list(csv.reader(io.StringIO(text)))[:30]
        return {"kind": "csv", "rows": rows}
    return {"kind": "txt", "text": text[:4000]}
