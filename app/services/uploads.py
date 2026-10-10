"""Validated upload into the trusted store. Document text cannot change controls."""

from __future__ import annotations

import uuid
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit.service import write_audit
from app.config import get_settings
from app.domain.dates import utcnow
from app.domain.errors import AccessDeniedError, IngestionError
from app.ingestion.pipeline import ingest_document
from app.ingestion.store import TrustedDocumentStore
from app.ingestion.validate import validate_upload
from app.models.tables import ControlState, Customer, Document
from app.security.principal import Principal


def accept_upload(
    session: Session,
    principal: Principal,
    *,
    filename: str,
    data: bytes,
    customer_id: str,
    risk_domain: str,
    document_type: str,
    reporting_date: date,
    extracted_hint: str | None = None,
) -> dict:
    if not principal.can("document.upload"):
        raise AccessDeniedError("This role cannot upload documents.")
    if principal.read_only:
        raise AccessDeniedError("This role is read-only.")
    customer = session.get(Customer, customer_id)
    from app.authorisation.service import allowed_customer_ids, deny

    if customer is None or customer.customer_id not in allowed_customer_ids(session, principal):
        deny(
            session,
            principal,
            customer_id=customer_id,
            detail="Upload blocked before the file was stored.",
            correlation_id=uuid.uuid4().hex,
        )
    settings = get_settings()
    demo = session.get(ControlState, "demo_mode")
    demo_mode = bool(demo and demo.value == "true")
    result = validate_upload(
        filename,
        data,
        max_bytes=settings.upload_max_bytes,
        allow_macros=settings.allow_macros,
        demo_mode=demo_mode,
        extracted_text=extracted_hint if extracted_hint is not None else _quick_text(filename, data),
    )
    document_id = f"UPL-{uuid.uuid4().hex[:12]}"
    if not result.accepted:
        session.add(
            Document(
                document_id=document_id,
                controlled_reference=f"HMB-QUAR-{document_id}",
                title=filename,
                document_type=document_type,
                original_filename=filename,
                safe_filename=result.safe_filename or "rejected.bin",
                storage_key=f"quarantine/{document_id}",
                mime_type=result.mime_type or "application/octet-stream",
                checksum_sha256=result.checksum or uuid.uuid4().hex,
                file_size=len(data),
                group_id=customer.group_id,
                portfolio_id=customer.portfolio_id,
                enterprise=False,
                risk_domain=risk_domain,
                authoring_function="uploader",
                reporting_date=reporting_date,
                uploaded_at=utcnow(),
                uploader=principal.username,
                version=1,
                superseded=False,
                confidentiality="Confidential",
                synthetic_flag=False,
                ingestion_status="quarantined",
                malware_scan_status="not_scanned_hackathon_placeholder",
                parser_status="rejected",
                source_viewer_route=f"/source?document_id={document_id}",
                retention_classification="hackathon-ephemeral-90-day-demonstration",
                legal_hold=False,
                sensitivity="standard",
                quarantine_reason=result.reason,
            )
        )
        write_audit(
            session,
            event_type="upload_quarantined",
            outcome="quarantined",
            detail=result.reason or "Rejected.",
            principal=principal,
            customer_id=customer_id,
            document_id=document_id,
        )
        raise IngestionError(result.reason or "The file was quarantined.")
    existing = session.scalar(select(Document).where(Document.checksum_sha256 == result.checksum))
    if existing is not None:
        raise IngestionError("A document with the same checksum is already stored.")
    key = f"uploads/{customer_id}/{result.safe_filename}"
    TrustedDocumentStore(settings.object_store_dir).put(key, data)
    document = Document(
        document_id=document_id,
        controlled_reference=f"HMB-UPL-{document_id}",
        title=filename,
        document_type=document_type,
        original_filename=filename,
        safe_filename=result.safe_filename or filename,
        storage_key=key,
        mime_type=result.mime_type or "application/octet-stream",
        checksum_sha256=result.checksum or "",
        file_size=len(data),
        group_id=customer.group_id,
        portfolio_id=customer.portfolio_id,
        enterprise=False,
        risk_domain=risk_domain,
        authoring_function="uploader",
        reporting_date=reporting_date,
        uploaded_at=utcnow(),
        uploader=principal.username,
        version=1,
        superseded=False,
        confidentiality="Confidential",
        synthetic_flag=True,
        ingestion_status="pending",
        malware_scan_status="not_scanned_hackathon_placeholder",
        parser_status="pending",
        source_viewer_route=f"/source?document_id={document_id}",
        retention_classification="hackathon-ephemeral-90-day-demonstration",
        legal_hold=False,
        sensitivity="standard",
    )
    session.add(document)
    session.flush()
    from app.models.tables import DocumentCustomer

    session.add(DocumentCustomer(document_id=document_id, customer_id=customer_id))
    status = ingest_document(
        session,
        TrustedDocumentStore(settings.object_store_dir),
        document,
        correlation_id=uuid.uuid4().hex,
        actor=principal.username,
    )
    write_audit(
        session,
        event_type="upload",
        outcome=status,
        detail="Upload processed by the ingestion pipeline.",
        principal=principal,
        customer_id=customer_id,
        document_id=document_id,
    )
    return {
        "document_id": document_id,
        "ingestion_status": document.ingestion_status,
        "parser_status": document.parser_status,
        "quarantine_reason": document.quarantine_reason,
        "prompt_injection_suspected": document.prompt_injection_suspected,
    }


def _quick_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith((".txt", ".csv")):
        return data.decode("utf-8", errors="replace")
    if lower.endswith(".pdf"):
        try:
            import fitz

            pdf = fitz.open(stream=data, filetype="pdf")
            return "\n".join(page.get_text() for page in pdf)
        except Exception:
            return ""
    if lower.endswith(".docx"):
        try:
            from docx import Document as Docx

            return "\n".join(p.text for p in Docx(io_bytes(data)).paragraphs)
        except Exception:
            return ""
    if lower.endswith(".xlsx"):
        try:
            import io

            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), data_only=True)
            parts = []
            for name in book.sheetnames:
                for row in book[name].iter_rows(max_row=20, max_col=6, values_only=True):
                    parts.append(" ".join("" if cell is None else str(cell) for cell in row))
            return "\n".join(parts)
        except Exception:
            return ""
    return ""


def io_bytes(data: bytes):
    import io

    return io.BytesIO(data)
