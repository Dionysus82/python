"""Word extraction that keeps heading path, paragraph number and tables."""

from __future__ import annotations

import io

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.parsers.base import ParsedChunk, ParsedDocument, mark_injection


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        return ParsedDocument(text="", corrupt=True, error=str(exc))
    chunks: list[ParsedChunk] = []
    heading_path: list[str] = []
    paragraph_number = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            level = 1
            tail = style.replace("Heading", "").strip()
            if tail.isdigit():
                level = int(tail)
            heading_path = heading_path[: level - 1]
            heading_path.append(text or style)
        if not text:
            continue
        paragraph_number += 1
        chunks.append(
            ParsedChunk(
                text=text,
                locator={
                    "paragraph": paragraph_number,
                    "heading_path": " > ".join(heading_path),
                    "section": heading_path[-1] if heading_path else None,
                },
            )
        )
    table_meta = []
    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(rows).strip()
        table_meta.append({"table_number": table_number, "rows": len(table.rows)})
        if text:
            chunks.append(
                ParsedChunk(
                    text=text,
                    locator={
                        "table_number": table_number,
                        "section": heading_path[-1] if heading_path else "Table",
                        "heading_path": " > ".join(heading_path),
                    },
                )
            )
    parsed = ParsedDocument(
        text="\n".join(chunk.text for chunk in chunks),
        chunks=chunks,
        section_metadata={"headings": heading_path, "tables": table_meta},
    )
    mark_injection(parsed.chunks)
    return parsed
