"""Excel extraction that retains sheet, table and cell-range metadata."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.parsers.base import ParsedChunk, ParsedDocument, mark_injection


def _column(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", corrupt=True, error=str(exc))
    chunks: list[ParsedChunk] = []
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append(
            {
                "name": sheet.title,
                "tables": list(sheet.tables.keys()),
                "dimensions": sheet.dimensions,
            }
        )
        rows = list(sheet.iter_rows())
        if not rows:
            continue
        header_map: dict[int, str] = {}
        for cell in rows[0]:
            if cell.value is not None:
                header_map[cell.column] = str(cell.value)
        for row in rows:
            values = []
            numeric_cells = []
            for cell in row:
                if cell.value is None:
                    continue
                header = header_map.get(cell.column, _column(cell.column))
                kind = "formula" if getattr(cell, "data_type", None) == "f" else "value"
                values.append(f"{header}={cell.value} ({kind})")
                numeric_cells.append(f"{_column(cell.column)}{cell.row}")
            if not values:
                continue
            first = numeric_cells[0]
            last = numeric_cells[-1]
            chunks.append(
                ParsedChunk(
                    text=f"Sheet {sheet.title} row {row[0].row}: " + " | ".join(values),
                    locator={
                        "sheet": sheet.title,
                        "cell_range": f"{first}:{last}" if first != last else first,
                        "row": row[0].row,
                        "tables": list(sheet.tables.keys()),
                    },
                )
            )
    parsed = ParsedDocument(
        text="\n".join(chunk.text for chunk in chunks),
        chunks=chunks,
        sheet_metadata={"sheets": sheets},
    )
    mark_injection(parsed.chunks)
    workbook.close()
    return parsed
